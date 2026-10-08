import os
import sys
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC = Document()
sec = DOC.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

style = DOC.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
for h in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
    st = DOC.styles[h]
    st.font.name = 'Times New Roman'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')


def add_image(path, maxw_cm=15.5):
    if not os.path.exists(path):
        print('missing image', path)
        return
    w, h = Image.open(path).size
    width = Cm(maxw_cm)
    if w / h > 4:
        width = Cm(maxw_cm)
    DOC.add_picture(path, width=width)
    DOC.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    src, out = sys.argv[1], sys.argv[2]
    base = os.path.dirname(os.path.abspath(src))
    in_code = False
    buf = []
    for raw in open(src, encoding='utf-8').read().split('\n'):
        line = raw.rstrip('\r')
        if line.startswith('```'):
            if in_code:
                p = DOC.add_paragraph()
                for l in buf:
                    r = p.add_run(l + '\n')
                    r.font.name = 'Consolas'
                    r.font.size = Pt(8.5)
                    r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                buf = []
            in_code = not in_code
            continue
        if in_code:
            buf.append(line)
            continue
        if not line.strip():
            continue
        if line.startswith('!['):
            path = line[line.find('](') + 2:line.rfind(')')]
            add_image(os.path.join(base, path))
        elif line.strip() == '[[PAGEBREAK]]':
            DOC.add_page_break()
        elif line.startswith('%%#'):
            p = DOC.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[3:])
            r.font.size = Pt(18)
            r.bold = True
            r.font.name = 'Times New Roman'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        elif line.startswith('%%'):
            p = DOC.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.font.size = Pt(14)
            r.font.name = 'Times New Roman'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        elif line.startswith('#### '):
            DOC.add_heading(line[5:], level=4)
        elif line.startswith('### '):
            DOC.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            DOC.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            DOC.add_heading(line[2:], level=1)
        elif line.startswith('- '):
            DOC.add_paragraph(line[2:], style='List Bullet')
        else:
            DOC.add_paragraph(line)
    DOC.save(out)
    print('saved', out)


main()